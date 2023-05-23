# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import pandas as pd
import os
import time
import sys
from lib.ppt_word_checker_ui import Ui_MainWindow
from lib.logger import Logger
from PyQt5.QtWidgets import *
from PyQt5 import QtCore


class Main(QMainWindow, Ui_MainWindow):
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.locate = os.getcwd()
        self.logging = Logger()
        self.logging.log.info("Complete !!")
        self.locate = os.getcwd()
        self.file_locate = os.listdir(self.locate + "\\inspection")
        self.worktype_runs =[]
        self.filename_runs = []
        self.page_runs = []
        self.type_runs = []
        self.findword_runs = []
        self.text_runs = []
        self.oldtext_runs = []
        self.newtext_runs = []
        self.find_texts = []
        self.change_texts = {}
        self.find_text_load()
        self.change_text_load()

        self.plainTextEdit.appendPlainText("**********************************************************************"
                                           "\n PPT 내 텍스트를 찾고, 수정할 수 있습니다."
                                           "\n   Collect All -  폴더 내 파일의 모든 텍스트를 수집합니다. "
                                           "\n   Find All - 폴더 내 파일의 특정 텍스트를 수집합니다. "
                                           "\n   Change All - 폴더 내 파일의 특정 텍스트를 다른 텍스트로 수정합니다."
                                           "\n**********************************************************************"
                                           "\n 검색 요청 단어는 총 " + str(len(self.find_texts)) + "건입니다."
                                           "\n 변경 요청 단어는 총 " + str(len(self.change_texts)) + "건입니다."
                                           "\n**********************************************************************"
                                           "\n 검색 요청 리스트 \n"
                                            + str(self.find_texts) +
                                           "\n**********************************************************************"
                                           "\n 변경 요청 리스트 \n"
                                            + str(self.change_texts) +
                                           "\n**********************************************************************"
                                           "\n\n\n")
        # 1. Collect All 버튼 클릭
        self.pushButton_2.clicked.connect(self.pushButton_2_clicked)
        # 2. Find All 버튼 클릭
        self.pushButton_3.clicked.connect(self.pushButton_3_clicked)
        # 3. Change All 버튼 클릭
        self.pushButton_5.clicked.connect(self.pushButton_5_clicked)
        self.logging.log.info("Complete !!")

    def find_text_load(self):
        self.wordlist_find = self.locate + "\\WordList\\Find_Word_List.xlsx"
        self.df_find = pd.read_excel(self.wordlist_find, usecols="A")
        self.df_find_value = self.df_find.values.tolist()
        self.logging.log.info("Complete !!")
        self.i = 0
        while self.i < len(self.df_find_value):
            self.find_texts.append(self.df_find_value[self.i][0])
            self.i += 1
            self.logging.log.info("Complete !!")

    def change_text_load(self):
        self.wordlist_find = self.locate + "\\WordList\\Change_word_List.xlsx"
        self.df_change = pd.read_excel(self.wordlist_find, usecols="A:B")
        self.df_change_value = self.df_change.values.tolist()
        self.logging.log.info("Complete !!")
        self.i = 0
        while self.i < len(self.df_change_value):
            self.change_texts[self.df_change_value[self.i][0]] = self.df_change_value[self.i][1]
            # print(self.change_texts)
            self.i += 1
            self.logging.log.info("Complete !!")

    def data_append_all(self, worktype, filename, page, shape_type, findword, oldword, newword, text):
        self.worktype_runs.append(worktype)
        self.filename_runs.append(filename)
        self.page_runs.append("page " + str(page))
        self.type_runs.append(shape_type)
        self.text_runs.append(text)
        self.findword_runs.append(findword)
        self.oldtext_runs.append(oldword)
        self.newtext_runs.append(newword)
        self.logging.log.info("Complete !!")

    def extract_excel(self):
        self.df = pd.DataFrame()
        self.df["work_type"] = self.worktype_runs
        self.df["File_name"] = self.filename_runs
        self.df["page_number"] = self.page_runs
        self.df["Type"] = self.type_runs
        self.df["Find_word"] = self.findword_runs
        self.df["old_text"] = self.oldtext_runs
        self.df["new_text"] = self.newtext_runs
        self.df["Text_value"] = self.text_runs
        self.df1 = self.df[["work_type","File_name", "page_number", "Type","Find_word","old_text", "new_text", "Text_value"]]
        self.cur_time = time.strftime('%Y%m%d_%H%M%S', time.localtime(time.time()))
        # self.df1.to_excel(self.locate + "\\ppt_word_check_list_"+ self.cur_time +".xlsx", encoding="EUC-KR")
        self.df1.to_excel(self.locate + "\\ppt_word_check_list_" + self.cur_time + ".xlsx" )
        self.df1.to_csv(self.locate + "\\ppt_word_check_list_"+ self.cur_time +".csv", encoding="utf-8")
        self.logging.log.info("Complete !!")

    def runs_reset(self):
        self.worktype_runs = []
        self.filename_runs = []
        self.page_runs = []
        self.type_runs = []
        self.findword_runs = []
        self.text_runs = []
        self.oldtext_runs = []
        self.newtext_runs = []

    # 1. Collect All 버튼 클릭
    def pushButton_2_clicked(self):
        QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
        self.plainTextEdit.appendPlainText("Collect All 작업을 시작합니다.")
        for(path, dir, file) in os.walk(str(self.locate + "\\inspection")):
            print("path:" + path)
            try:
                for i in os.listdir(path):
                    if i[-4:] != "pptx":
                        continue
                    prs = Presentation(str(path+"\\"+i))
                    self.file_name = i
                    QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
                    self.plainTextEdit.appendPlainText(str(i) + "파일을 찾았습니다. 작업을 시작합니다.")
                    self.page_count = 0
                    for slide in prs.slides:
                        self.page_count += 1
                        self.logging.log.info("Complete !!")
                        # case1
                        for shape in slide.shapes:
                            # print(shape)
                            if shape.has_text_frame:
                                for shape_paragraph in shape.text_frame.paragraphs:
                                    self.data_append_all("Collect", i, self.page_count, "shape",None, None, None, shape_paragraph.text)

                            if shape.has_table:
                                for table_paragraph in shape.table.iter_cells():
                                    self.data_append_all("Collect", i, self.page_count, "table",None, None, None, table_paragraph.text)

                            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                                for group_paragraph in shape.shapes:
                                    if group_paragraph.has_text_frame:
                                        self.data_append_all("Collect", i, self.page_count, "group",None, None, None,group_paragraph.text)
            except:
                pass
        self.extract_excel()
        self.runs_reset()
        QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
        self.plainTextEdit.appendPlainText("Collect All 작업을 종료하였습니다.\n")

    # 2. Find All 버튼 클릭
    def pushButton_3_clicked(self):
        QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
        self.plainTextEdit.appendPlainText("Find All 작업을 시작합니다.")
        for (path, dir, file) in os.walk(str(self.locate + "\\inspection")):
            print("path:" + path)
            for i in os.listdir(path):
                if i[-4:] != "pptx":
                    continue
                print(str(i) + "시작")
                prs = Presentation(str(path+"\\"+ str(i)))
                self.file_name = i
                self.page_count = 0
                self.logging.log.info("Complete !!")
                QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
                self.plainTextEdit.appendPlainText(str(i) + "파일을 찾았습니다. 작업을 시작합니다.")
                for slide in prs.slides:
                    self.page_count += 1
                    for shape in slide.shapes:
                        try:
                            if shape.has_text_frame:
                                for shape_paragraph in shape.text_frame.paragraphs:
                                    for search_text in self.find_texts:
                                        if search_text in shape_paragraph.text:
                                            self.data_append_all("Find", i, self.page_count, "shape", search_text, None, None, shape_paragraph.text)
                                            print(shape_paragraph.text)
                        except:
                            pass
                        try:
                            if shape.has_table:
                                for table_paragraph in shape.table.iter_cells():
                                    for search_text in self.find_texts:
                                        if search_text in table_paragraph.text:
                                            self.data_append_all("Find", i, self.page_count, "table", search_text, None, None, table_paragraph.text)
                                            print(table_paragraph.text)
                        except:
                            pass
                        try:
                            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                                for group_paragraph in shape.shapes:
                                    if group_paragraph.has_text_frame:
                                        for search_text in self.find_texts:
                                            if search_text in group_paragraph.text:
                                                self.data_append_all("Find", i, self.page_count, "group", search_text, None, None, group_paragraph.text)
                                                print(group_paragraph.text)
                        except:
                            pass
                print(str(i) + "종료")
                QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
                self.plainTextEdit.appendPlainText(str(i) + "파일에 대한 작업을 완료하였습니다.\n")

        self.extract_excel()
        self.runs_reset()
        QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
        self.plainTextEdit.appendPlainText("Find All 작업을 종료하였습니다.\n")

    # 3. Change All 버튼 클릭
    def pushButton_5_clicked(self):
        QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
        self.plainTextEdit.appendPlainText("Change All 작업을 시작합니다.")
        for (path, dir, file) in os.walk(str(self.locate + "\\inspection")):
            print("path:" + path)
            for i in os.listdir(path):
                if i[-4:] != "pptx":
                    continue
                prs = Presentation(str(path+"\\"+i))
                self.file_name = i
                self.page_count = 0
                self.logging.log.info("Complete !!")
                QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
                self.plainTextEdit.appendPlainText(str(i) + "파일을 찾았습니다. 작업을 시작합니다.")
                for slide in prs.slides:
                    self.page_count += 1

                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for shape_paragraph in shape.text_frame.paragraphs:
                                for search_text in self.change_texts.keys():
                                    for run in shape_paragraph.runs:
                                        cur_text = run.text
                                        if search_text in cur_text:
                                            self.data_append_all("Change", i, self.page_count, "shape", None, search_text, self.change_texts[search_text], cur_text)
                                            # print(cur_text + "//" + search_text)
                                        new_text = cur_text.replace(search_text,self.change_texts[search_text])
                                        run.text = new_text

                        if shape.has_table:
                            for row in shape.table.rows:
                                for cell in row.cells:
                                    for table_paragraph in cell.text_frame.paragraphs:
                                        for search_text in self.change_texts.keys():
                                            for run in table_paragraph.runs:
                                                cur_text = run.text
                                                if search_text in cur_text:
                                                    self.data_append_all("Change", i, self.page_count, "shape", None, search_text, self.change_texts[search_text], cur_text)
                                                    # print(cur_text + "//" + search_text)
                                                new_text = cur_text.replace(search_text, self.change_texts[search_text])
                                                run.text = new_text


                        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                            for group_paragraph in shape.shapes:
                                if group_paragraph.has_text_frame:
                                    for paragraph in group_paragraph.text_frame.paragraphs:
                                        for search_text in self.change_texts.keys():
                                            for run in paragraph.runs:
                                                cur_text = run.text
                                                if search_text in cur_text:
                                                    self.data_append_all("Change", i, self.page_count, "shape", None, search_text, self.change_texts[search_text], cur_text)
                                                    # print(cur_text + "//" + search_text)
                                                new_text = cur_text.replace(search_text, self.change_texts[search_text])
                                                run.text = new_text
                QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
                self.plainTextEdit.appendPlainText(str(i) + "파일에 대한 작업을 완료하였습니다.\n")
                prs.save(self.locate + "\\inspection\\" + str(i))
        self.extract_excel()
        self.runs_reset()
        QtCore.QCoreApplication.processEvents()  # 아래 메시지 생성이 딜레이 되는 것을 방지
        self.plainTextEdit.appendPlainText("Change All 작업을 종료하였습니다.\n")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    you_view_main = Main()
    you_view_main.show()
    app.exec()